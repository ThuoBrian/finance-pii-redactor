"""Tests for the Excel write adapter: highlighting plus the Crosswalk sheet.

Framework-free besides pandas/openpyxl themselves: no Presidio/spaCy involved,
so these run without the language model (matching the rest of the suite).
"""

from __future__ import annotations

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from finance_redactor.domain.errors import UnreadableFileError
from finance_redactor.infrastructure.documents.excel_gateway import (
    OpenpyxlExcelGateway,
)

_CROSSWALK_COLUMNS = [
    "Original name",
    "Entity type",
    "Category",
    "Pseudonym",
    "Flagged",
    "Possible match",
]


def _crosswalk_df(rows: list[dict]) -> pd.DataFrame:
    return pd.DataFrame(rows, columns=_CROSSWALK_COLUMNS)


def test_read_rejects_a_file_that_is_not_really_a_workbook():
    """A corrupted or wrongly-renamed file surfaces as the shared error type,
    not openpyxl's own ``BadZipFile``/``KeyError``.
    """
    with pytest.raises(UnreadableFileError):
        OpenpyxlExcelGateway().read_sheets(BytesIO(b"not a workbook file at all"))


def test_write_produces_redacted_and_crosswalk_sheets():
    df = pd.DataFrame({"notes": ["Paid to STF-10010"]})
    crosswalk_df = _crosswalk_df(
        [
            {
                "Original name": "Jane Doe",
                "Entity type": "PERSON",
                "Category": "Staff",
                "Pseudonym": "STF-10010",
                "Flagged": "",
                "Possible match": "",
            }
        ]
    )

    result = OpenpyxlExcelGateway().write([("Redacted", df, set())], crosswalk_df)

    workbook = load_workbook(BytesIO(result))
    assert workbook.sheetnames == ["Redacted", "Crosswalk"]


def test_redacted_sheet_content_and_highlighting_are_preserved():
    df = pd.DataFrame({"notes": ["Paid to STF-10010", "no name here"]})

    result = OpenpyxlExcelGateway().write(
        [("Redacted", df, {(0, "notes")})], _crosswalk_df([])
    )

    workbook = load_workbook(BytesIO(result))
    sheet = workbook["Redacted"]
    # Header row, then two data rows.
    assert sheet.cell(row=1, column=1).value == "notes"
    assert sheet.cell(row=2, column=1).value == "Paid to STF-10010"
    assert sheet.cell(row=3, column=1).value == "no name here"
    # Row 0 ("notes" -> column 1) was flagged as changed; row 1 was not.
    assert sheet.cell(row=2, column=1).fill.fgColor.rgb == "00FFFF00"
    assert sheet.cell(row=3, column=1).fill.fgColor.rgb in (None, "00000000")


def test_crosswalk_sheet_matches_the_passed_dataframe():
    df = pd.DataFrame({"notes": ["x"]})
    crosswalk_df = _crosswalk_df(
        [
            {
                "Original name": "Jane Doe",
                "Entity type": "PERSON",
                "Category": "Staff",
                "Pseudonym": "STF-10010",
                "Flagged": "",
                "Possible match": "",
            },
            {
                "Original name": "Micheal Sample",
                "Entity type": "PERSON",
                "Category": "",
                "Pseudonym": "PSN-AUTO-0BA3D",
                "Flagged": "yes",
                "Possible match": "Michael Sample (STF-12345, 92% match)",
            },
        ]
    )

    result = OpenpyxlExcelGateway().write([("Redacted", df, set())], crosswalk_df)

    read_back = pd.read_excel(BytesIO(result), sheet_name="Crosswalk")
    assert list(read_back.columns) == _CROSSWALK_COLUMNS
    assert read_back.iloc[0]["Original name"] == "Jane Doe"
    assert read_back.iloc[1]["Flagged"] == "yes"
    # Header row is bolded for readability.
    sheet = load_workbook(BytesIO(result))["Crosswalk"]
    assert sheet.cell(row=1, column=1).font.bold is True


def test_write_handles_an_empty_crosswalk_without_crashing():
    df = pd.DataFrame({"notes": ["no PII here"]})

    result = OpenpyxlExcelGateway().write([("Redacted", df, set())], _crosswalk_df([]))

    read_back = pd.read_excel(BytesIO(result), sheet_name="Crosswalk")
    assert list(read_back.columns) == _CROSSWALK_COLUMNS
    assert len(read_back) == 0


def _round_tripped_workbook_bytes(data: dict) -> bytes:
    """Write ``data`` to real xlsx bytes and back, like an actual upload.

    Constructing a DataFrame directly in memory (``pd.DataFrame({...})``)
    does not reproduce the dtype a real upload gets: pandas 3.0 infers its
    own string dtype when *reading* an xlsx file, which can differ from
    whatever dtype an in-memory literal happens to get. Round-tripping
    through real bytes via ``gateway.read()`` is what actually exercises
    ``text_columns()`` the way the app does.
    """
    buffer = BytesIO()
    pd.DataFrame(data).to_excel(buffer, index=False)
    buffer.seek(0)
    return buffer.getvalue()


def test_text_columns_includes_string_columns_from_a_real_upload():
    # Regression test: pandas 3.0 infers a dedicated string dtype (shown as
    # "str" in df.dtypes) for text columns read from an xlsx file, not the
    # legacy "object" dtype - a `dtype == object` check silently matched
    # nothing, so no column was ever pre-selected for a real uploaded file.
    gateway = OpenpyxlExcelGateway()
    df = gateway.read_sheets(
        BytesIO(
            _round_tripped_workbook_bytes(
                {"Description": ["Consulting payment"], "Payee": ["Jane Doe"]}
            )
        )
    )["Sheet1"]

    assert set(gateway.text_columns(df)) == {"Description", "Payee"}


def test_text_columns_excludes_numeric_and_date_columns_from_a_real_upload():
    gateway = OpenpyxlExcelGateway()
    df = gateway.read_sheets(
        BytesIO(
            _round_tripped_workbook_bytes(
                {
                    "Payee": ["Jane Doe"],
                    "Amount": [1000],
                    "Date": pd.to_datetime(["2020-01-01"]),
                }
            )
        )
    )["Sheet1"]

    assert gateway.text_columns(df) == ["Payee"]


def test_hyperlink_targets_never_reach_the_output():
    """Output is rebuilt from values, so a cell's link target is dropped.

    Guards against a future switch to editing the workbook in place, which
    would silently carry ``mailto:`` targets through.
    """
    from openpyxl import Workbook

    source = Workbook()
    cell = source.active["A2"]
    source.active["A1"] = "notes"
    cell.value = "Contact"
    cell.hyperlink = "mailto:jane.doe@example.org"
    buffer = BytesIO()
    source.save(buffer)
    buffer.seek(0)

    gateway = OpenpyxlExcelGateway()
    ((name, df),) = gateway.read_sheets(buffer).items()
    output = gateway.write([(name, df, set())], _crosswalk_df([]))

    workbook = load_workbook(BytesIO(output))
    assert workbook[name]["A2"].value == "Contact"
    assert all(
        not sheet_cell.hyperlink
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for sheet_cell in row
    )


def _workbook_bytes(
    tabs: dict[str, pd.DataFrame], hidden: tuple[str, ...] = ()
) -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, df in tabs.items():
            df.to_excel(writer, index=False, sheet_name=name)
    workbook = load_workbook(BytesIO(buffer.getvalue()))
    for name in hidden:
        workbook[name].sheet_state = "hidden"
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def test_read_sheets_returns_every_tab_in_order_including_hidden_ones():
    data = _workbook_bytes(
        {
            "First": pd.DataFrame({"a": ["x"]}),
            "Second": pd.DataFrame({"a": ["y"]}),
            "Third": pd.DataFrame({"a": ["z"]}),
        },
        hidden=("Second",),
    )

    sheets = OpenpyxlExcelGateway().read_sheets(BytesIO(data))

    assert list(sheets) == ["First", "Second", "Third"]
    assert sheets["Second"].loc[0, "a"] == "y"


def test_write_keeps_every_tab_and_highlights_the_right_cells_per_tab():
    first = pd.DataFrame({"notes": ["one", "two"]})
    second = pd.DataFrame({"notes": ["three", "four"]})

    result = OpenpyxlExcelGateway().write(
        [("Jan", first, {(0, "notes")}), ("Feb", second, {(1, "notes")})],
        _crosswalk_df([]),
    )

    workbook = load_workbook(BytesIO(result))
    assert workbook.sheetnames == ["Jan", "Feb", "Crosswalk"]
    assert workbook["Jan"]["A2"].fill.fgColor.rgb == "00FFFF00"
    assert workbook["Jan"]["A3"].fill.fgColor.rgb in (None, "00000000")
    assert workbook["Feb"]["A2"].fill.fgColor.rgb in (None, "00000000")
    assert workbook["Feb"]["A3"].fill.fgColor.rgb == "00FFFF00"


def test_write_makes_tab_names_valid_and_unique_without_dropping_a_tab():
    df = pd.DataFrame({"notes": ["x"]})
    long_name = "L" * 40
    tabs = [
        ("VND-10011", df, set()),
        ("VND-10011", df, set()),  # two tabs redacted to the same name
        ("[ACCOUNT]", df, set()),  # a mask holds characters Excel forbids
        (long_name, df, set()),
        ("Crosswalk", df, set()),  # collides with the tool's own sheet
        ("", df, set()),
    ]

    result = OpenpyxlExcelGateway().write(tabs, _crosswalk_df([]))

    names = load_workbook(BytesIO(result)).sheetnames
    assert len(names) == len(tabs) + 1
    assert len({n.lower() for n in names}) == len(names)
    assert names[-1] == "Crosswalk"
    assert "VND-10011" in names
    assert "VND-10011 (2)" in names
    assert "_ACCOUNT_" in names
    assert long_name[:31] in names
    assert "Crosswalk (2)" in names
    assert "Sheet6" in names
    assert all(len(n) <= 31 for n in names)
