"""Tests for the python-docx gateway.

Uses real ``python-docx`` objects (no mocks) to build fixture documents, the
same style as ``test_pdf_gateway.py`` for PyMuPDF.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from docx.enum.section import WD_SECTION

from finance_redactor.domain.entities import Span
from finance_redactor.domain.errors import UnreadableFileError
from finance_redactor.infrastructure.documents.docx_gateway import PythonDocxDocument


def _build_docx(builder) -> bytes:
    """Build a .docx via ``builder(document)`` and return its bytes."""
    document = Document()
    builder(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_open_rejects_a_file_that_is_not_really_a_docx():
    """A corrupted or wrongly-renamed file surfaces as the shared error type,
    not python-docx's own ``BadZipFile``/``KeyError``.
    """
    with pytest.raises(UnreadableFileError):
        PythonDocxDocument.open(b"not a docx file at all")


def test_blocks_cover_body_table_and_header_paragraphs():
    def build(document):
        document.add_paragraph("A memo about John Doe.")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Vendor: Acme Supplies"
        document.sections[0].header.paragraphs[0].text = "Confidential memo"

    gateway = PythonDocxDocument.open(_build_docx(build))
    texts = [gateway.block_text(i) for i in range(gateway.block_count)]

    assert "A memo about John Doe." in texts
    assert "Vendor: Acme Supplies" in texts
    assert "Confidential memo" in texts


def test_linked_headers_are_not_scanned_twice():
    def build(document):
        document.add_paragraph("body one")
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph("body two")
        document.sections[0].header.paragraphs[0].text = "Shared header"

    gateway = PythonDocxDocument.open(_build_docx(build))
    texts = [gateway.block_text(i) for i in range(gateway.block_count)]

    assert texts.count("Shared header") == 1


def test_replace_block_text_confined_to_one_run():
    def build(document):
        document.add_paragraph("Hello World and more text")

    gateway = PythonDocxDocument.open(_build_docx(build))
    text = gateway.block_text(0)
    start = text.index("World")

    gateway.replace_block_text(0, [(Span(start, start + len("World")), "STF-1")])

    assert gateway.block_text(0) == "Hello STF-1 and more text"


def test_replace_block_text_preserves_formatting_outside_the_span():
    def build(document):
        paragraph = document.add_paragraph("Hello ")
        bold_run = paragraph.add_run("World")
        bold_run.bold = True
        paragraph.add_run(" and more text")

    gateway = PythonDocxDocument.open(_build_docx(build))
    text = gateway.block_text(0)
    start = text.index("World")

    # Span crosses from the plain prefix run into the bold "World" run.
    gateway.replace_block_text(0, [(Span(start - 2, start + len("World")), "PSN-2")])

    runs = gateway._paragraphs[0].runs
    assert runs[-1].text == " and more text"
    assert runs[-1].bold is None
    assert "PSN-2" in "".join(r.text for r in runs)


def test_replace_block_text_handles_multiple_spans_in_one_call():
    def build(document):
        document.add_paragraph("John paid Mary today")

    gateway = PythonDocxDocument.open(_build_docx(build))
    text = gateway.block_text(0)
    john = text.index("John")
    mary = text.index("Mary")

    gateway.replace_block_text(
        0,
        [
            (Span(john, john + 4), "STF-1"),
            (Span(mary, mary + 4), "STF-2"),
        ],
    )

    assert gateway.block_text(0) == "STF-1 paid STF-2 today"


def test_to_bytes_round_trips_replacements():
    def build(document):
        document.add_paragraph("Invoice for John Doe")

    gateway = PythonDocxDocument.open(_build_docx(build))
    text = gateway.block_text(0)
    start = text.index("John Doe")
    gateway.replace_block_text(0, [(Span(start, start + len("John Doe")), "STF-99")])

    reopened = PythonDocxDocument.open(gateway.to_bytes())
    assert "STF-99" in reopened.block_text(0)
    assert "John" not in reopened.block_text(0)


# --- Hyperlinks ---------------------------------------------------------------

_TARGET = "mailto:jane.doe@example.org"


def _add_hyperlink(paragraph, text: str, target: str = _TARGET) -> None:
    """Append a real external ``<w:hyperlink>`` (python-docx has no API for it)."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    link.append(run)
    paragraph._p.append(link)


def _add_field_run(paragraph, *, fld_char: str | None = None, instr: str | None = None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = OxmlElement("w:r")
    if fld_char:
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), fld_char)
        run.append(element)
    if instr:
        element = OxmlElement("w:instrText")
        element.text = instr
        run.append(element)
    paragraph._p.append(run)


def test_hyperlink_text_is_scanned_and_its_target_is_removed():
    """Regression: ``paragraph.runs`` skipped link runs, so the text leaked too."""

    def build(document):
        paragraph = document.add_paragraph("Contact ")
        _add_hyperlink(paragraph, "Jane Doe")

    gateway = PythonDocxDocument.open(_build_docx(build))

    assert gateway.block_text(0) == "Contact Jane Doe"
    assert gateway.remove_external_links() == 1
    gateway.replace_block_text(0, [(Span(8, 16), "STF-10010")])
    output = gateway.to_bytes()
    assert b"mailto:" not in output
    reopened = Document(BytesIO(output))
    assert reopened.paragraphs[0].text == "Contact STF-10010"
    assert all("hyperlink" not in rel.reltype for rel in reopened.part.rels.values())


def test_hyperlink_in_a_header_is_removed():
    def build(document):
        _add_hyperlink(document.sections[0].header.paragraphs[0], "Jane Doe")

    gateway = PythonDocxDocument.open(_build_docx(build))

    assert gateway.remove_external_links() == 1
    assert "Jane Doe" in [gateway.block_text(i) for i in range(gateway.block_count)]
    assert b"mailto:" not in gateway.to_bytes()


def test_hyperlink_field_code_is_removed_and_its_result_kept():
    def build(document):
        paragraph = document.add_paragraph("See ")
        _add_field_run(paragraph, fld_char="begin")
        _add_field_run(paragraph, instr=f' HYPERLINK "{_TARGET}" ')
        _add_field_run(paragraph, fld_char="separate")
        paragraph.add_run("Jane Doe")
        _add_field_run(paragraph, fld_char="end")

    gateway = PythonDocxDocument.open(_build_docx(build))

    assert gateway.remove_external_links() == 1
    assert gateway.block_text(0) == "See Jane Doe"
    output = gateway.to_bytes()
    assert b"mailto:" not in output
    assert b"fldChar" not in output


def test_simple_hyperlink_field_is_unwrapped():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def build(document):
        paragraph = document.add_paragraph("See ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), f' HYPERLINK "{_TARGET}" ')
        run = OxmlElement("w:r")
        text_el = OxmlElement("w:t")
        text_el.text = "Jane Doe"
        run.append(text_el)
        field.append(run)
        paragraph._p.append(field)

    gateway = PythonDocxDocument.open(_build_docx(build))

    assert gateway.remove_external_links() == 1
    assert gateway.block_text(0) == "See Jane Doe"
    assert b"mailto:" not in gateway.to_bytes()


def test_document_without_links_reports_zero():
    gateway = PythonDocxDocument.open(
        _build_docx(lambda d: d.add_paragraph("No links here."))
    )
    assert gateway.remove_external_links() == 0
