"""Tests for carrying a workbook's look (not its content) into the redacted copy.

Real openpyxl workbooks with made-up data. The redacted copy is rebuilt from
values, so the point of these tests is twofold: the look survives, and nothing
else (comments, formulas, validation lists, defined names) does.
"""

from __future__ import annotations

import datetime
import zipfile
from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from finance_redactor.infrastructure.documents.excel_formatting import (
    safe_number_format,
)
from finance_redactor.infrastructure.documents.excel_gateway import (
    OpenpyxlExcelGateway,
)

_EMPTY_CROSSWALK = pd.DataFrame(columns=["Original name"])


def _formatted_workbook() -> bytes:
    """Build a made-up report: title row, styles, and risky extras."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "Quarterly payments"
    ws.merge_cells("A1:C1")
    for cell, text in (
        ("A3", "Payee"),
        ("B3", "Amount"),
        ("C3", "Date"),
        ("D3", "Total"),
    ):
        ws[cell] = text
        ws[cell].font = Font(bold=True)
        ws[cell].fill = PatternFill("solid", fgColor="DDEEFF")
    ws["A4"] = "Someone"
    ws["B4"] = 1234.5
    ws["B4"].number_format = '"KES" #,##0.00'
    ws["C4"] = datetime.date(2026, 3, 1)
    ws["C4"].number_format = "dd-mmm-yy"
    ws["D4"] = "=B4*2"
    ws["B5"] = 99
    ws["B5"].number_format = '"Paid to a named person" #,##0'
    ws["A4"].comment = Comment("Ask Pat Example", "Sam Author")
    ws["A4"].hyperlink = "mailto:someone@example.org"
    ws.column_dimensions["A"].width = 40
    ws.row_dimensions[3].height = 30
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = "FF0000"
    validation = DataValidation(type="list", formula1='"Alpha Vendor,Beta Vendor"')
    ws.add_data_validation(validation)
    validation.add("A5")
    ws.conditional_formatting.add(
        "B4:B5",
        CellIsRule(
            operator="greaterThan",
            formula=["100"],
            fill=PatternFill("solid", fgColor="FF0000"),
        ),
    )
    wb.defined_names["HiddenNameRange"] = DefinedName(
        "HiddenNameRange", attr_text="Report!$A$4"
    )
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _round_trip(data: bytes, tab: str = "Report") -> bytes:
    gateway = OpenpyxlExcelGateway()
    sheets = gateway.read_sheets(BytesIO(data))
    formats = gateway.read_formatting(BytesIO(data))
    return gateway.write([(tab, sheets[tab], set())], _EMPTY_CROSSWALK, [formats[tab]])


def test_the_look_of_a_title_row_sheet_survives():
    out = load_workbook(BytesIO(_round_trip(_formatted_workbook())))["Report"]

    assert out["A1"].value == "Quarterly payments"
    assert "A1:C1" in {str(r) for r in out.merged_cells.ranges}
    assert out["A3"].value == "Payee"  # same position as in the source
    assert out["A3"].font.bold is True
    assert out["A3"].fill.fgColor.rgb == "00DDEEFF"
    assert out["B4"].number_format == '"KES" #,##0.00'
    assert out["C4"].number_format == "dd-mmm-yy"
    assert out.column_dimensions["A"].width == 40
    assert out.row_dimensions[3].height == 30
    assert out.freeze_panes == "A4"
    assert out.sheet_properties.tabColor.rgb == "00FF0000"


def test_a_blank_header_cell_is_not_written_as_unnamed():
    out = load_workbook(BytesIO(_round_trip(_formatted_workbook())))["Report"]

    assert out["B1"].value is None  # pandas called it "Unnamed: 1"
    assert out["D1"].value is None


def test_comments_formulas_validation_names_and_links_do_not_come_through():
    output = _round_trip(_formatted_workbook())

    with zipfile.ZipFile(BytesIO(output)) as archive:
        everything = b"".join(archive.read(name) for name in archive.namelist())
    for leaked in (
        b"Ask Pat Example",
        b"Sam Author",
        b"Alpha Vendor",
        b"HiddenNameRange",
        b"B4*2",
        b"mailto:",
        b"Paid to a named person",
    ):
        assert leaked not in everything, leaked


def test_what_was_left_behind_is_counted_for_the_note():
    formats = OpenpyxlExcelGateway().read_formatting(BytesIO(_formatted_workbook()))

    assert formats["Report"].dropped == {
        "formulas": 1,
        "comments": 1,
        "validation or conditional-formatting rules": 2,
    }


def test_a_number_format_with_readable_text_falls_back_to_general():
    out = load_workbook(BytesIO(_round_trip(_formatted_workbook())))["Report"]

    assert out["B5"].number_format == "General"


def test_safe_number_format_keeps_ordinary_formats_and_drops_text():
    assert safe_number_format("0.00%") == "0.00%"
    assert safe_number_format('"KES" #,##0.00') == '"KES" #,##0.00'
    assert safe_number_format("[$USD-409]#,##0") == "[$USD-409]#,##0"
    assert safe_number_format('"Paid to Jane Doe" @') == "General"
    assert safe_number_format(r"\P\a\i\d\ \t\o\ \J\a\n\e") == "General"
    assert safe_number_format("0" * 200) == "General"


def test_the_changed_cell_highlight_wins_over_the_source_fill():
    wb = Workbook()
    ws = wb.active
    ws.title = "T"
    ws["A1"] = "notes"
    ws["A2"] = "one"
    ws["A3"] = "two"
    ws["A2"].fill = PatternFill("solid", fgColor="FF0000")
    ws["A3"].fill = PatternFill("solid", fgColor="FF0000")
    buffer = BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()
    gateway = OpenpyxlExcelGateway()
    sheets = gateway.read_sheets(BytesIO(data))
    formats = gateway.read_formatting(BytesIO(data))

    out = load_workbook(
        BytesIO(
            gateway.write(
                [("T", sheets["T"], {(0, "notes")})], _EMPTY_CROSSWALK, [formats["T"]]
            )
        )
    )["T"]

    assert out["A2"].fill.fgColor.rgb == "00FFFF00"  # changed: yellow
    assert out["A3"].fill.fgColor.rgb == "00FF0000"  # untouched: source red


def test_each_tab_gets_its_own_formatting_whatever_it_is_renamed_to():
    wb = Workbook()
    first = wb.active
    first.title = "One"
    first["A1"] = "a"
    first.column_dimensions["A"].width = 11
    second = wb.create_sheet("Two")
    second["A1"] = "b"
    second.column_dimensions["A"].width = 33
    buffer = BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()
    gateway = OpenpyxlExcelGateway()
    sheets = gateway.read_sheets(BytesIO(data))
    formats = gateway.read_formatting(BytesIO(data))

    out = load_workbook(
        BytesIO(
            gateway.write(
                [
                    ("VND-10011", sheets["One"], set()),
                    ("ORG-20022", sheets["Two"], set()),
                ],
                _EMPTY_CROSSWALK,
                [formats["One"], formats["Two"]],
            )
        )
    )

    assert out["VND-10011"].column_dimensions["A"].width == 11
    assert out["ORG-20022"].column_dimensions["A"].width == 33


def test_unreadable_formatting_gives_a_plain_workbook_not_an_error():
    assert OpenpyxlExcelGateway().read_formatting(BytesIO(b"not a workbook")) == {}

    out = OpenpyxlExcelGateway().write(
        [("T", pd.DataFrame({"a": ["x"]}), set())], _EMPTY_CROSSWALK, [None]
    )

    assert load_workbook(BytesIO(out)).sheetnames == ["T", "Crosswalk"]
