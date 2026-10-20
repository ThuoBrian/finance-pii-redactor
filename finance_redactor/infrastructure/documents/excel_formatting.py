"""Carry a workbook's look (not its content) from the source to the redacted copy.

The redacted workbook is rebuilt from cell values on purpose: nothing from the
source except cell text can reach it. This module keeps that property while
still preserving how the sheet looks, by copying only an allow-list of
presentation settings. Anything not listed here is dropped, so a feature added
to Excel later cannot leak by being forgotten.

Copied: column widths and hidden state, row heights and hidden state, per cell
font / fill / border / alignment / number format, merged ranges, freeze panes,
tab colour, zoom, gridlines and the auto-filter range.

Not copied, and counted so the user is told: formulas (the value is kept),
comments, images and charts, data validation and conditional-formatting rules.
Comments and formulas can hold names, and rules can hold literal text, so they
stay out. Hyperlinks never come through either.

Positions line up with the DataFrame because ``pd.read_excel`` reads from sheet
row 1 and column A with row 1 as the header, so DataFrame row ``i`` is sheet row
``i + 2`` and DataFrame column ``j`` is sheet column ``j + 1``.
"""

from __future__ import annotations

import re
from copy import copy
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# A number format may hold literal text (``"Paid to Jane Doe" @``). Only short
# currency-style literals are kept; anything longer falls back to General.
_MAX_LITERAL = 4
_MAX_FORMAT_LENGTH = 100
_QUOTED = re.compile(r'"([^"]*)"')
_ESCAPED_CHAR = re.compile(r"\\.")

_Style = tuple[Any, Any, Any, Any, str]  # font, fill, border, alignment, number format


@dataclass(frozen=True)
class SheetFormat:
    """The allow-listed look of one sheet, plus what was left behind."""

    # (first column, last column, width, hidden) per column-dimension group.
    columns: tuple[tuple[int, int, float | None, bool], ...] = ()
    rows: dict[int, tuple[float | None, bool]] = field(default_factory=dict)
    styles: tuple[_Style, ...] = ()
    cell_styles: dict[tuple[int, int], int] = field(default_factory=dict)
    merged: tuple[str, ...] = ()
    freeze_panes: str | None = None
    tab_color: str | None = None
    zoom: int | None = None
    show_gridlines: bool | None = None
    auto_filter: str | None = None
    dropped: dict[str, int] = field(default_factory=dict)


def safe_number_format(number_format: str) -> str:
    """Return ``number_format``, or ``General`` if it could carry readable text."""
    if len(number_format) > _MAX_FORMAT_LENGTH:
        return "General"
    if any(len(lit) > _MAX_LITERAL for lit in _QUOTED.findall(number_format)):
        return "General"
    # A backslash escapes one literal character, so a run of them spells text.
    if len(_ESCAPED_CHAR.findall(number_format)) > _MAX_LITERAL:
        return "General"
    return number_format


def capture(workbook: Workbook) -> dict[str, SheetFormat]:
    """Collect the allow-listed look of every sheet, keyed by tab name."""
    return {sheet.title: _capture_sheet(sheet) for sheet in workbook.worksheets}


def _capture_sheet(sheet: Worksheet) -> SheetFormat:
    unique: dict[_Style, int] = {}
    cell_styles: dict[tuple[int, int], int] = {}
    formulas = comments = 0
    # ``_cells`` holds only cells that really exist in the file. ``iter_rows``
    # would materialize the whole bounding rectangle, which for a sheet with a
    # stray format at the far corner means millions of empty cells.
    for cell in list(sheet._cells.values()):
        if getattr(cell, "data_type", "") == "f":
            formulas += 1
        if getattr(cell, "comment", None) is not None:
            comments += 1
        if not cell.has_style or cell.row is None or cell.column is None:
            continue
        style = (
            copy(cell.font),
            copy(cell.fill),
            copy(cell.border),
            copy(cell.alignment),
            safe_number_format(cell.number_format),
        )
        cell_styles[(cell.row, cell.column)] = unique.setdefault(style, len(unique))

    tab = sheet.sheet_properties.tabColor
    tab_rgb = tab.rgb if tab is not None and tab.type == "rgb" else None
    view = sheet.sheet_view
    dropped = {
        "formulas": formulas,
        "comments": comments,
        "images or charts": len(getattr(sheet, "_images", []))
        + len(getattr(sheet, "_charts", [])),
        "validation or conditional-formatting rules": len(
            sheet.data_validations.dataValidation
        )
        + len(sheet.conditional_formatting),
    }
    return SheetFormat(
        columns=tuple(
            (dim.min or 1, dim.max or dim.min or 1, dim.width, bool(dim.hidden))
            for dim in sheet.column_dimensions.values()
            if dim.min
        ),
        rows={
            index: (dim.height, bool(dim.hidden))
            for index, dim in sheet.row_dimensions.items()
            if dim.height is not None or dim.hidden
        },
        styles=tuple(unique),
        cell_styles=cell_styles,
        merged=tuple(str(rng) for rng in sheet.merged_cells.ranges),
        freeze_panes=sheet.freeze_panes,
        tab_color=tab_rgb if isinstance(tab_rgb, str) else None,
        zoom=view.zoomScale,
        show_gridlines=view.showGridLines,
        auto_filter=sheet.auto_filter.ref,
        dropped={name: count for name, count in dropped.items() if count},
    )


def apply(sheet: Worksheet, fmt: SheetFormat, n_rows: int, n_cols: int) -> None:
    """Apply ``fmt`` to a rebuilt sheet of ``n_rows`` data rows and ``n_cols`` columns.

    Row 1 is the header, so the sheet is ``n_rows + 1`` rows tall.
    """
    last_row = n_rows + 1
    for first, last, width, hidden in fmt.columns:
        for index in range(first, min(last, n_cols) + 1):
            dimension = sheet.column_dimensions[get_column_letter(index)]
            if width is not None:
                dimension.width = width
            dimension.hidden = hidden
    for index, (height, hidden) in fmt.rows.items():
        if index <= last_row:
            sheet.row_dimensions[index].height = height
            sheet.row_dimensions[index].hidden = hidden
    for (row, column), style_index in fmt.cell_styles.items():
        if row > last_row or column > n_cols:
            continue
        font, fill, border, alignment, number_format = fmt.styles[style_index]
        cell = sheet.cell(row=row, column=column)
        cell.font = copy(font)
        cell.fill = copy(fill)
        cell.border = copy(border)
        cell.alignment = copy(alignment)
        cell.number_format = number_format
    for merged_range in fmt.merged:
        sheet.merge_cells(merged_range)
    if fmt.freeze_panes:
        sheet.freeze_panes = fmt.freeze_panes
    if fmt.tab_color:
        sheet.sheet_properties.tabColor = fmt.tab_color
    if fmt.zoom:
        sheet.sheet_view.zoomScale = fmt.zoom
    if fmt.show_gridlines is not None:
        sheet.sheet_view.showGridLines = fmt.show_gridlines
    if fmt.auto_filter:
        sheet.auto_filter.ref = fmt.auto_filter
